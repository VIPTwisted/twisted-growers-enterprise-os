"""Manifest Bridge API — FastAPI wrapper over the copied slc/planner engine.

Run from the backend folder (any cwd works — deps.py pins it):
    .venv\\Scripts\\python -m uvicorn app.main:app --port 8010 --reload
"""

import io
import os
import shutil
import time
from datetime import datetime
from typing import Optional

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response
from pydantic import BaseModel

from . import ai, blast, bridge, deps, inventory, labtools, metrctools, picksheet, review, tasklog

slc = deps.slc
S = deps.session_state

app = FastAPI(title="Manifest Bridge API")


@app.on_event("startup")
def _start_background_warm():
    # All periodic network refreshes (received checks, destination backfill,
    # active pull, bulk order re-warm) run on this daemon thread — /api/board
    # never does inline network work.
    bridge.start_background_warm()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:3000", "http://localhost:3001"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
    # the .xlsx export is fetched via POST, so JS has to read the filename the
    # server picked — Content-Disposition isn't CORS-safelisted by default
    expose_headers=["Content-Disposition"],
)


# ── models ───────────────────────────────────────────────────
class LinkBody(BaseModel):
    digits: str
    manifest_id: str


class GroupBody(BaseModel):
    manifest_id: str
    digits: list[str]


class VerifyBody(BaseModel):
    digits: Optional[str] = None


class QuickShipBody(BaseModel):
    digits: str
    day: str = "today"          # today | tomorrow
    terms: str = "net30"        # net30 | cod


class OpenBody(BaseModel):
    digits: str
    force_live: bool = False


class QtyBody(BaseModel):
    uid: str
    qty: int                     # negative = relative adjustment
    handle: str = "inventory"    # inventory | trash


class RemoveBody(BaseModel):
    uids: list[str]
    disposition: str = "return"  # return | delete


class UndoBody(BaseModel):
    index: int


class AddBody(BaseModel):
    product_id: int | str
    batch_id: int | str
    qty: int
    unit: str = "Unit"           # Unit | Case


class PaymentBody(BaseModel):
    amount: float
    memo: str = ""
    accounting_method: str = "payment"
    payment_type: str = ""
    payment_date: Optional[str] = None


class PaymentDeleteBody(BaseModel):
    payment_id: int | str


class CreditAdjustBody(BaseModel):
    amount: float
    memo: str = ""


class MirrorBody(BaseModel):
    payment_id: int | str
    sibling_digits: str


class MetaBody(BaseModel):
    delivery_date: Optional[str] = None
    due_date: Optional[str] = None
    net_terms_name: Optional[str] = None
    stage_id: Optional[int | str] = None


class TierBody(BaseModel):
    tier_id: Optional[int | str] = None


class PairApplyBody(BaseModel):
    sibling_digits: str


class SplitSel(BaseModel):
    id: int | str
    qty: int


class SplitBody(BaseModel):
    selections: list[SplitSel]
    target_digits: Optional[str] = None


class ContactMember(BaseModel):
    name: str = ""
    email: str = ""


class ContactsBody(BaseModel):
    lists: dict[str, list[ContactMember]]


class ContactsXlsxBody(BaseModel):
    """A rectangular block lifted straight off the contacts grid — whatever the
    user highlighted, including edits that haven't been saved yet."""
    headers: list[str] = []
    rows: list[list[str]] = []
    store_col: int = -1  # index within `headers` holding the vendor, -1 = none


class BlastGroup(BaseModel):
    store: str = ""
    emails: list[str] = []


class BlastAttachment(BaseModel):
    name: str = ""
    mime: str = "application/octet-stream"
    b64: str = ""


class BlastBody(BaseModel):
    groups: list[BlastGroup]
    subject: str = ""
    body: str = ""
    cc: list[str] = []
    attachments: list[BlastAttachment] = []
    gap_seconds: float = 2.0


class CandQuery(BaseModel):
    batch: str
    item: str = ""
    grams: float = 0


class CandBatchBody(BaseModel):
    queries: list[CandQuery]


class PickSheetBody(BaseModel):
    digits: list[str]


class EmailBody(BaseModel):
    to: list[str]
    cc: list[str] = []
    subject: str
    body: str
    attach: str = "pdf"          # pdf | bundle
    extra: list[BlastAttachment] = []   # user-picked files (e.g. a back-order sheet)


class AiParseBody(BaseModel):
    text: str


class AiChatBody(BaseModel):
    messages: list[dict]


class BatchSetBody(BaseModel):
    batch_id: int | str
    product_id: int | str
    quantity: Optional[int] = None
    name: Optional[str] = None


class AiSendBody(BaseModel):
    to: list[str]
    cc: list[str] = []
    subject: str
    body: str


# ── helpers ──────────────────────────────────────────────────
def _board_ctx():
    orders = bridge.load_pack_orders()
    manifests = bridge.load_manifests()
    links = bridge.load_links()
    views = [bridge.manifest_view(m) for m in manifests]
    suggestions = bridge.auto_match(orders, manifests, links, views=views)
    checks = bridge.load_verify_cache()
    edges = bridge.resolve(orders, manifests, links, suggestions, checks,
                           views=views)
    return orders, manifests, edges, checks


def _find_pair(digits, orders, manifests, edges):
    o = next((x for x in orders if x["digits"] == digits), None)
    e = next((x for x in edges if x["invoice"] == digits), None)
    m = next((x for x in manifests
              if str(x.get("Id")) == (e or {}).get("manifest_id")), None)
    return o, m, e


def _review_response():
    orders, manifests, edges, _checks = _board_ctx()
    payload = review.review_payload(orders, edges, manifests)
    if payload is None:
        raise HTTPException(404, "No pair loaded — open a review first.")
    return payload


# ── data sources (the old slc sidebar) ───────────────────────
# CSV index reads are heavy (full read_csv) — memoized 5 min, exactly like
# slc's own @st.cache_data(ttl=300) which the shim strips.
_CSV_STATS = {"at": 0.0, "data": None}


def _csv_stats():
    if _CSV_STATS["data"] is not None and time.time() - _CSV_STATS["at"] < 300:
        return _CSV_STATS["data"]
    try:
        csv_data, _src = slc.load_csv_index()
    except Exception:
        csv_data = []
    csv_data = csv_data or []
    last_date, _p = slc.get_csv_last_date()
    out = {
        "transfers": len(csv_data),
        "with_invoices": sum(1 for t in csv_data if t.get("invoice_number")),
        "latest": last_date,
    }
    _CSV_STATS.update(at=time.time(), data=out)
    return out


@app.get("/api/datasources")
def datasources():
    cache = S.get("mp_cache") or slc.load_manifest_cache()
    S["mp_cache"] = cache
    manifests = cache.get("manifests", [])
    try:
        slc.enrich_recipients_from_index(manifests)
    except Exception:
        pass
    missing = sum(1 for t in manifests if slc.recipient_of(t)[0] in ("", "—"))
    inv = slc.load_inventory_cache()
    binv = slc.load_bapi_inventory()
    return {
        "csv": _csv_stats(),
        "manifest_cache": {
            "count": len(manifests),
            "pulled_at": cache.get("pulled_at"),
            "window_days": slc.MP_WINDOW_DAYS,
            "in_window": len(slc.manifests_in_window(manifests)),
            "missing_dest": missing,
        },
        "contact_lists": len(slc._load_contact_lists() or {}),
        "apex_inventory": {"products": len(inv.get("products") or []),
                           "pulled_at": inv.get("pulled_at")},
        "bapi_inventory": {"batches": len(binv.get("rows") or []),
                           "pulled_at": binv.get("pulled_at")},
    }


@app.post("/api/metrc/full-pull")
def full_pull():
    cache, changes = slc.refresh_full(S.get("mp_cache") or slc.load_manifest_cache())
    slc.save_manifest_cache(cache)
    S["mp_cache"] = cache
    ch = changes or {}
    return {"ok": True,
            "message": (f"Full {slc.MP_WINDOW_DAYS}-day pull · "
                        f"{ch.get('total', len(cache.get('manifests', [])))} cached · "
                        f"{len(ch.get('new') or [])} new · "
                        f"{len(ch.get('changed') or [])} changed.")}


@app.post("/api/csv/update")
def csv_update():
    try:
        n = slc.incremental_csv_update()
    except Exception as ex:
        return {"ok": False, "message": f"CSV update failed: {type(ex).__name__}: {ex}"}
    _CSV_STATS["at"] = 0.0
    slc.get_csv_last_date.clear()
    return {"ok": True, "message": f"📂 CSV index updated — {n or 0} new row(s) merged."}


@app.post("/api/inventory/load")
def inventory_load():
    """⚠️ METERED — Apex v1 products pull (find-all). User-triggered only."""
    cache, changes = slc.refresh_inventory_full(slc.load_inventory_cache())
    slc.save_inventory_cache(cache)
    S["inv_cache"] = cache
    return {"ok": True,
            "message": (f"📦 Apex inventory loaded — {len(cache.get('products') or [])} "
                        f"products (metered v1 call).")}


@app.post("/api/inventory/refresh")
def inventory_refresh():
    """⚠️ METERED — Apex v1 products changed since the last pull. User-triggered only."""
    cache, changes = slc.refresh_inventory_changes(slc.load_inventory_cache())
    slc.save_inventory_cache(cache)
    S["inv_cache"] = cache
    ch = changes or {}
    return {"ok": True,
            "message": (f"📦 Inventory refresh — {len(ch.get('new') or [])} new · "
                        f"{len(ch.get('changed') or [])} changed (metered v1 call).")}


@app.post("/api/inventory/bapi-pull")
def inventory_bapi_pull():
    sess = S.get("bapi_session") or {}
    if not slc._session_ready(sess):
        return {"ok": False,
                "message": "Batch inventory pull needs the 🔑 b-api session — grab it first."}
    rows, msg = slc.bapi_fetch_inventory(sess, show_progress=False)
    if rows:
        from datetime import datetime as _dt
        binv = {"pulled_at": _dt.now().isoformat(timespec="seconds"), "rows": rows}
        slc.save_bapi_inventory(binv)
        S["bapi_inv"] = binv
        # Fresh b-api rows: rebuild the batch index now (it may be serving the
        # b-api fallback while the bearer API is 429-capped) and drop cached ➕
        # candidates so /api/review/candidates re-matches against the new batches.
        slc.fetch_all_batches_indexed.clear()
        for k in [k for k in S if str(k).startswith("_cmp_cands_")]:
            S.pop(k, None)
    return {"ok": bool(rows), "message": f"🧾 {msg} (🔑 b-api, 0 metered calls)"}


# ── board ────────────────────────────────────────────────────
@app.get("/api/board")
def get_board(received_check: bool = True):
    return bridge.board_payload(run_received_check=received_check)


@app.get("/api/prefetch/status")
def get_prefetch():
    return bridge.prefetch_status()


@app.post("/api/metrc/quick-pull")
def quick_pull():
    cache, changes = slc.refresh_quick(S.get("mp_cache") or slc.load_manifest_cache())
    slc.save_manifest_cache(cache)
    S["mp_cache"] = cache
    ids, _at = bridge.refresh_active_ids()
    _log_metrc_new((changes or {}).get("new"))
    return {"ok": True,
            "message": (f"Pulled active outgoing · "
                        f"{len((changes or {}).get('new') or [])} new · "
                        f"{len(ids)} manifest(s) currently in transit.")}


def _log_metrc_new(new_list):
    """METRC refresh scoring: +1 per INVOICE affected by the freshly pulled
    manifests — 10 new manifests tagged to 2 invoices score 2, because two
    invoices were updated. A manifest with no invoice tag scores once itself."""
    invs, untagged = {}, []
    for m in (new_list or []):
        if not isinstance(m, dict):
            untagged.append(str(m)[:40])
            continue
        mn = str(m.get("ManifestNumber") or m.get("Id") or "")[:40]
        digs = bridge.manifest_invoices(m)
        if digs:
            for dg in digs:
                invs.setdefault(dg, []).append(mn)
        else:
            untagged.append(mn)
    for dg, mns in invs.items():
        tasklog.log_task("metrc_manifest_new", invoice=f"Twiste-{dg}",
                         detail="manifest " + ", ".join(mns[:3]))
    tasklog.log_many("metrc_manifest_new", untagged)


@app.post("/api/metrc/incremental-pull")
def incremental_pull():
    cache, changes = slc.refresh_incremental(
        S.get("mp_cache") or slc.load_manifest_cache())
    slc.save_manifest_cache(cache)
    S["mp_cache"] = cache
    ch = changes or {}
    _log_metrc_new(ch.get("new"))
    return {"ok": True,
            "message": (f"Incremental pull · {len(ch.get('new') or [])} new · "
                        f"{len(ch.get('changed') or [])} changed.")}


@app.post("/api/metrc/fill-destinations")
def fill_destinations():
    cache = S.get("mp_cache") or slc.load_manifest_cache()
    filled, cache = slc.fill_destinations_via_deliveries(cache)
    slc.save_manifest_cache(cache)
    S["mp_cache"] = cache
    return {"ok": True, "message": f"Filled {filled} store name(s)."}


@app.post("/api/apex/refresh-orders")
def refresh_orders():
    try:
        ok, msg = bridge.refresh_apex_orders()
    finally:
        deps.progress(active=False)
    return {"ok": ok, "message": msg}


# ── live task progress (polled by the frontend's top bar) ────
@app.get("/api/progress")
def get_progress():
    p = S.get("_task_progress") or {}
    return p if p.get("active") else {"active": False}


# ── links ────────────────────────────────────────────────────
@app.post("/api/links")
def set_link(body: LinkBody):
    bridge.set_link(body.digits, body.manifest_id, mode="manual")
    return {"ok": True, "message": f"🔗 Linked Twiste-{body.digits}."}


@app.delete("/api/links/{digits}")
def clear_link(digits: str):
    bridge.clear_link(digits)
    return {"ok": True, "message": f"✂️ Unlinked Twiste-{digits}."}


@app.post("/api/links/group")
def group_links(body: GroupBody):
    links_now = bridge.load_links()
    have = {dg for dg, rec in links_now.items()
            if str(rec.get("manifest_id")) == str(body.manifest_id)}
    want = set(body.digits)
    for dg in want - have:
        bridge.set_link(dg, body.manifest_id, mode="manual")
    for dg in have - want:
        rec = links_now.get(dg)
        if rec and str(rec.get("manifest_id")) == str(body.manifest_id):
            bridge.clear_link(dg)
    return {"ok": True,
            "message": (f"🔗 {len(want)} invoice(s) linked"
                        + (" — they'll compare COMBINED." if len(want) > 1 else "."))}


@app.post("/api/links/clear")
def clear_all_links():
    bridge.save_links({})
    return {"ok": True, "message": "All saved links cleared."}


# ── content verification ─────────────────────────────────────
@app.post("/api/verify")
def verify(body: VerifyBody):
    orders, manifests, edges, checks = _board_ctx()
    if body.digits:
        todo = [e for e in edges if e["invoice"] == body.digits]
    else:
        todo = [e for e in edges if e["state"] != "mismatch"]
    done = 0
    fresh_compared = []
    try:
        for i, e in enumerate(todo):
            o = next((x for x in orders if x["digits"] == e["invoice"]), None)
            m = next((x for x in manifests
                      if str(x.get("Id")) == e["manifest_id"]), None)
            if not o or not m:
                continue
            deps.progress(f"⚖️ Comparing {o['invoice']} "
                          f"({i + 1}/{len(todo)})…",
                          int(100 * i / max(len(todo), 1)))
            sibs = bridge.group_siblings(o, m, orders, edges)
            try:
                res = bridge.verify_pair(o, m, sibs)
                checks[bridge.verify_key(o["digits"], e["manifest_id"])] = res
                # Make.com-style metering: +1 per line item scanned & confirmed
                # on a FRESH compare — shared scorer with prefetch/review-open.
                fresh_compared.append((e.get("state"), res, o["invoice"]))
            except Exception as ex:
                checks[bridge.verify_key(o["digits"], e["manifest_id"])] = {
                    "state": "error", "detail": f"{type(ex).__name__}: {ex}"}
            done += 1
        bridge.save_verify_cache(checks)
        for prev_state, res, inv in fresh_compared:
            bridge.score_fresh_compare(prev_state, res, inv)
    finally:
        deps.progress(active=False)
    return {"ok": True, "message": f"Compared {done} pair(s)."}


@app.post("/api/verify/clear")
def clear_verify():
    bridge.save_verify_cache({})
    return {"ok": True, "message": "Comparison results cleared."}


# ── ⚡ quick ship ─────────────────────────────────────────────
@app.post("/api/quick-ship")
def quick_ship(body: QuickShipBody):
    orders = bridge.load_pack_orders()
    o = next((x for x in orders if x["digits"] == body.digits), None)
    if not o:
        return {"ok": False, "message": "⚡ That pack order is no longer on the board."}
    ok, msg = bridge.quick_ship(o, body.day, body.terms)
    return {"ok": ok, "message": msg}


# ── 🔑 session ────────────────────────────────────────────────
@app.get("/api/session")
def get_session():
    sess = S.get("bapi_session") or {}
    ok = slc._session_ready(sess)
    return {"loaded": ok, "summary": slc._session_summary(sess) if ok else None}


@app.post("/api/session/grab")
def grab_session():
    sess = S.get("bapi_session") or {}
    cid = sess.get("company_id") or slc.DEFAULT_COMPANY_ID
    s, authed, msg = slc.firefox_session_all_in_one(cid)
    if s:
        S["bapi_session"] = s
        deps.save_session(s)
    if authed:
        return {"ok": True, "message": f"🔑 Session valid ✓ — {msg}"}
    if s:
        return {"ok": False, "message": f"⚠️ Grabbed a session but it didn't pass: {msg}"}
    return {"ok": False, "message": f"❌ {msg}"}


@app.post("/api/session/test")
def test_session():
    sess = S.get("bapi_session") or {}
    ok, code, msg = slc.bapi_test_session(sess)
    return {"ok": ok, "message": (f"🧪 Session valid ✓ — {msg}" if ok
                                  else f"🧪 Session FAILED — {msg}")}


# ── push log ─────────────────────────────────────────────────
@app.get("/api/pushlog")
def get_pushlog():
    return {"log": S.get("apex_push_log") or []}


@app.post("/api/pushlog/clear")
def clear_pushlog():
    S.pop("apex_push_log", None)
    S.pop("apex_push_verdict", None)
    return {"ok": True}


# ── 📥 deal-flow orders (Submitted / Accepted / Finalized) ───
@app.get("/api/orders/by-status")
def orders_by_status(live: bool = False):
    """Saved locally like every other cache: the default read serves the
    banked copy instantly (0 network); live=true (the ↻ button) re-pulls
    from Apex and re-banks. 🧾 Refresh orders also re-banks it."""
    if not live:
        cached = bridge.load_dealflow_cache()
        if cached:
            return {"ok": True, "note": "",
                    "submitted": cached.get("submitted") or [],
                    "accepted": cached.get("accepted") or [],
                    "finalized": cached.get("finalized") or [],
                    "at": cached.get("at") or "", "cached": True}
    ok, note, submitted, accepted, finalized = bridge.deal_flow_orders()
    return {"ok": ok, "note": note, "submitted": submitted,
            "accepted": accepted, "finalized": finalized,
            "at": datetime.now().isoformat(timespec="seconds"),
            "cached": False}


@app.get("/api/orders/stages")
def order_stages(digits: str):
    """Deal-flow stage list, read off one order's own deal_flow — used to
    populate the bulk stage picker."""
    sess = S.get("bapi_session") or {}
    if not slc._session_ready(sess):
        return {"stages": [], "note": "Needs the 🔑 b-api session."}
    _oid, od, err = bridge.fetch_apex_order(f"Twiste-{digits}", digits)
    if not od:
        return {"stages": [], "note": err or "Order not found."}
    sts = ((od.get("deal_flow") or {}).get("order_statuses")) or []
    return {"stages": [{"id": s.get("id"), "name": s.get("name")}
                       for s in sts if not s.get("archived")],
            "note": ""}


class BulkStageBody(BaseModel):
    digits: list[str]
    stage_id: Optional[str] = None
    stage_name: str = ""


@app.post("/api/orders/bulk-stage")
def orders_bulk_stage(body: BulkStageBody):
    """Set the deal-flow stage on MANY orders in one shot. Each order resolves
    the stage id off its own deal_flow (by id if given, else by name), then
    PATCHes exactly like the single-order editor."""
    sess = S.get("bapi_session") or {}
    if not slc._session_ready(sess):
        return {"ok": False, "message": "Needs the 🔑 b-api session — grab it up top."}
    if not body.digits:
        return {"ok": False, "message": "Check at least one order first."}
    oks, fails = [], []
    for i, dg in enumerate(body.digits):
        inv = f"Twiste-{dg}"
        deps.progress(f"🔀 Deal flow {inv} ({i + 1}/{len(body.digits)})…",
                      int(100 * i / len(body.digits)))
        try:
            od, err = bridge.fetch_dealflow_order(dg)
            if not od:
                fails.append(f"{inv} — {err or 'not found'}")
                continue
            sid = body.stage_id
            sts = ((od.get("deal_flow") or {}).get("order_statuses")) or []
            if sid is None and body.stage_name:
                sid = next((s.get("id") for s in sts
                            if str(s.get("name", "")).strip().lower()
                            == body.stage_name.strip().lower()), None)
            if sid is None:
                fails.append(f"{inv} — stage '{body.stage_name}' not on this "
                             f"order's deal flow")
                continue
            cur = (od.get("order_status_id")
                   or (od.get("order_status") or {}).get("id"))
            if str(cur) == str(sid):
                oks.append(inv)          # already there — nothing to push
                continue
            log = []
            ok, fresh = slc.bapi_update_order_fields(
                sess, od, {"order_status_id": sid}, log)
            bridge.push_log(log)
            if ok:
                fresh = fresh or {}
                if fresh.get("id"):
                    bridge._order_cache_store(dg, fresh, fresh.get("id"))
                oks.append(inv)
                tasklog.log_task("deal_flow_updated", invoice=inv,
                                 detail=str(body.stage_name or sid))
            else:
                fails.append(f"{inv} — Apex rejected the update")
        except Exception as ex:
            fails.append(f"{inv} — {type(ex).__name__}")
    deps.progress(active=False)
    msg = (f"🔀 Deal flow → '{body.stage_name or body.stage_id}': "
           f"{len(oks)}/{len(body.digits)} updated.")
    if fails:
        msg += " ❌ " + " · ".join(fails[:5]) + ("…" if len(fails) > 5 else "")
    return {"ok": not fails, "message": msg}


@app.post("/api/review/open-order")
def review_open_order(body: OpenBody):
    """Open the review panel for an order WITHOUT a manifest (Submitted /
    Accepted deal-flow orders): Order details (deal flow, terms, dates) and
    item editing work; the compare grid activates once a manifest is linked."""
    inv = f"Twiste-{body.digits}"
    try:
        ok, msg = bridge.load_order_only(inv, body.digits,
                                         force_live=body.force_live)
        if not ok:
            raise HTTPException(502, msg)
        return _review_response()
    finally:
        deps.progress(active=False)


# ── 🔎 review ─────────────────────────────────────────────────
@app.post("/api/review/open")
def review_open(body: OpenBody):
    orders, manifests, edges, _checks = _board_ctx()
    o, m, e = _find_pair(body.digits, orders, manifests, edges)
    if not o or not m:
        raise HTTPException(400, "That pack order has no manifest linked to it.")
    try:
        if body.force_live:
            deps.progress(f"📡 Re-pull {o['invoice']} — refreshing the "
                          f"manifest cache (1 METRC call)…", 15)
            S["_force_live"] = body.digits
            bridge.order_cache_pop(body.digits)
            S.pop("_apex_cached_at", None)
            try:
                mc, _ch = slc.refresh_quick(S.get("mp_cache") or slc.load_manifest_cache())
                slc.save_manifest_cache(mc)
                S["mp_cache"] = mc
            except Exception:
                pass
            S.pop("selected_invoice", None)
            S.pop("comparison_df", None)
            S.pop("_cmp_sig", None)
            S.pop("_pay_synced_for", None)
        else:
            deps.progress(f"🔎 Opening {o['invoice']} — loading order + "
                          f"manifest…", 20)
        ok, msg = bridge.load_pair(o, m, force_live=body.force_live)
        if not ok:
            raise HTTPException(502, msg)
        review.fold_open_comparison_into_checks(o["digits"], str(m.get("Id")))
        return _review_response()
    finally:
        deps.progress(active=False)


@app.get("/api/review")
def review_get():
    return _review_response()


@app.post("/api/review/close")
def review_close():
    review.close_review()
    return {"ok": True}


@app.post("/api/review/qty")
def review_qty(body: QtyBody):
    ok, msg = review.set_line_quantity(body.uid, body.qty, body.handle)
    return {"ok": ok, "message": msg}


@app.post("/api/review/remove")
def review_remove(body: RemoveBody):
    ok, msg = review.remove_lines(body.uids, body.disposition)
    return {"ok": ok, "message": msg}


@app.post("/api/review/undo-remove")
def review_undo(body: UndoBody):
    ok, msg = review.undo_removed(body.index)
    return {"ok": ok, "message": msg}


@app.get("/api/review/candidates")
def review_candidates(batch: str, item: str = "", grams: float = 0):
    cands = review.inventory_candidates(batch, item)
    for c in cands:
        c["suggest"] = review.suggest_add_qty(c["product_id"], c["batch_id"], grams)
    return {"candidates": cands}


@app.post("/api/review/candidates-batch")
def review_candidates_batch(body: CandBatchBody):
    """All NOT-IN-INVOICE rows resolved in ONE request — the compare grid used
    to fire one /candidates GET per missing row."""
    out = {}
    for q in body.queries:
        cands = review.inventory_candidates(q.batch, q.item)
        for c in cands:
            c["suggest"] = review.suggest_add_qty(c["product_id"], c["batch_id"],
                                                  q.grams)
        out[q.batch] = cands
    return {"results": out}


@app.post("/api/review/add")
def review_add(body: AddBody):
    ok, msg = review.add_line(body.product_id, body.batch_id, body.qty, body.unit)
    return {"ok": ok, "message": msg}


@app.post("/api/review/add-all")
def review_add_all():
    """Every ❌ NOT IN INVOICE batch → top candidate at the suggested qty,
    in one shot — the server-side equivalent of clicking every ➕ row."""
    ok, msg = review.add_all_missing()
    return {"ok": ok, "message": msg}


@app.get("/api/review/pairscan")
def review_pairscan():
    orders, manifests, edges, _checks = _board_ctx()
    dg = review.loaded_digits()
    if not dg:
        raise HTTPException(404, "No pair loaded.")
    e_cur = next((x for x in edges if x["invoice"] == dg), None)
    m_cur = next((x for x in manifests
                  if str(x.get("Id")) == (e_cur or {}).get("manifest_id")), None)
    sibs = (bridge.group_siblings({"digits": dg}, m_cur, orders, edges)
            if m_cur is not None else [])
    return review.pair_scan(sibs)


# ── 💵 payments ───────────────────────────────────────────────
@app.post("/api/payments/sync")
def payments_sync():
    ok, msg = review.sync_payments()
    return {"ok": ok, "message": msg}


@app.post("/api/payments/add")
def payments_add(body: PaymentBody):
    ok, msg = review.add_payment(body.amount, body.memo, body.accounting_method,
                                 body.payment_type, body.payment_date)
    return {"ok": ok, "message": msg}


@app.post("/api/payments/adjust-credit")
def payments_adjust_credit(body: CreditAdjustBody):
    ok, msg = review.adjust_buyer_credit(body.amount, body.memo)
    return {"ok": ok, "message": msg}


@app.post("/api/payments/delete")
def payments_delete(body: PaymentDeleteBody):
    ok, msg = review.delete_payment(body.payment_id)
    return {"ok": ok, "message": msg}


@app.post("/api/payments/mirror")
def payments_mirror(body: MirrorBody):
    order = S.get("apex_order_raw") or {}
    rec = next((p for p in (order.get("payments") or [])
                if str(p.get("id")) == str(body.payment_id)), None)
    if not rec:
        return {"ok": False, "message": "That payment record wasn't found."}
    orders = bridge.load_pack_orders()
    sib = next((x for x in orders if x["digits"] == body.sibling_digits), None)
    if not sib:
        return {"ok": False, "message": "Sibling order not on the board."}
    ok, msg = bridge.pair_mirror_payment(rec, sib)
    return {"ok": ok, "message": msg}


# ── 📝 order details ──────────────────────────────────────────
@app.post("/api/order/meta")
def order_meta(body: MetaBody):
    ok, msg = review.save_meta(body.delivery_date, body.due_date,
                               body.net_terms_name, body.stage_id)
    return {"ok": ok, "message": msg}


@app.get("/api/order/pricing-tiers")
def pricing_tiers():
    sess = S.get("bapi_session") or {}
    if not slc._session_ready(sess):
        return {"tiers": [], "note": "💲 Pricing tier needs the Apex b-api session."}
    # Cached for the server's lifetime, same as slc's _ptiers session key —
    # tiers change rarely and this renders on every Order-details open.
    tiers = S.get("_ptiers")
    if tiers is None:
        tiers = bridge.bapi_list_pricing_tiers(sess)
        if tiers:
            S["_ptiers"] = tiers
    return {"tiers": [{"id": t.get("id"), "name": t.get("name"),
                       "percent": t.get("percent"),
                       "direction": t.get("direction"), "type": t.get("type")}
                      for t in tiers]}


@app.post("/api/order/pricing-tier")
def apply_tier(body: TierBody):
    sess = S.get("bapi_session") or {}
    if not slc._session_ready(sess):
        return {"ok": False, "message": "💲 Needs the Apex b-api session — 🔑 up top."}
    cur_ord = S.get("apex_order_raw") or {}
    ok, msg = bridge.bapi_change_pricing_tier(sess, cur_ord, body.tier_id)
    if ok:
        tasklog.log_task("pricing_tier_added", invoice=S.get("selected_invoice"),
                         detail=str(body.tier_id or "cleared"))
    return {"ok": ok, "message": msg}


@app.post("/api/order/pair-apply")
def pair_apply(body: PairApplyBody):
    live_src = S.get("apex_order_raw") or {}
    orders = bridge.load_pack_orders()
    sib = next((x for x in orders if x["digits"] == body.sibling_digits), None)
    if not sib:
        return {"ok": False, "message": "Sibling order not on the board."}
    ok, msg = bridge.pair_copy_meta(live_src, sib)
    return {"ok": ok, "message": msg}


# ── ✂️ split ──────────────────────────────────────────────────
@app.get("/api/split/items")
def split_items():
    rows, err = review.split_source()
    if rows is None:
        return {"items": [], "note": err}
    return {"items": rows, "note": ""}


@app.get("/api/split/available")
def split_available_route(term: str = ""):
    """Apex's own split-target picker list — searchable, any buyer."""
    orders, note = review.split_available(term)
    return {"orders": orders, "note": note}


class SplitPickBody(BaseModel):
    uuid: str
    digits: str = ""


@app.post("/api/split/target-pick")
def split_target_pick_route(body: SplitPickBody):
    res, err = review.split_target_pick(body.uuid, body.digits)
    return {"target": res, "note": err}


@app.get("/api/split/target")
def split_target(invoice: str):
    res, err = review.split_target_lookup(invoice)
    if res is None:
        return {"target": None, "note": err}
    return {"target": res, "note": ""}


@app.post("/api/split")
def do_split(body: SplitBody):
    ok, new_inv, msg = review.do_split(
        [{"id": s.id, "qty": s.qty} for s in body.selections],
        body.target_digits)
    return {"ok": ok, "new_invoice": new_inv, "message": msg}


# ── 📇 vendor contact lists (the email auto-fill source) ─────
@app.get("/api/contacts")
def contacts_get():
    return {"lists": slc._load_contact_lists()}


@app.post("/api/contacts")
def contacts_save(body: ContactsBody):
    """REPLACE email_contact_lists.json with the grid's contents — same
    semantics as the old Streamlit contact grid (deleting a row deletes that
    vendor's list). A timestamped backup is written first."""
    out = {}
    for store, members in body.lists.items():
        store = str(store).strip()
        if not store:
            continue
        mem, seen = [], set()
        for m in members or []:
            raw = (m.email or "").strip()
            em = slc._label_to_email(raw) or (raw if "@" in raw else "")
            if em and em.lower() not in seen:
                mem.append({"name": (m.name or "").strip(), "email": em})
                seen.add(em.lower())
        if mem:
            out[store] = mem
    path = os.path.join(deps.DATA_DIR, slc.CONTACT_LISTS_FILE)
    try:
        if os.path.exists(path):
            shutil.copyfile(path, path.replace(
                ".json", f".backup_{datetime.now():%Y%m%d_%H%M%S}.json"))
    except Exception:
        pass
    slc._save_contact_lists(out)
    return {"ok": True,
            "message": (f"📇 Saved {len(out)} vendor list(s) · "
                        f"{sum(len(v) for v in out.values())} email(s). "
                        f"Auto-fill uses these on the next review open.")}


@app.post("/api/contacts/xlsx")
def contacts_xlsx(body: ContactsXlsxBody):
    """Highlighted grid block → a real .xlsx download.

    Sheet 1 ("Contacts") mirrors the grid exactly — same rows, same columns,
    same cell text — so it round-trips back through Ctrl+V. Sheet 2 ("Emails")
    flattens it to one row per address (vendor · name · email), which is the
    shape mail-merge and list imports actually want. Sheet 2 is skipped when
    the block has no email-bearing cells."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    headers = [str(h) for h in body.headers]
    rows = [[("" if c is None else str(c)) for c in r] for r in body.rows]
    width = max([len(headers)] + [len(r) for r in rows] + [1])
    headers += [""] * (width - len(headers))
    rows = [r + [""] * (width - len(r)) for r in rows]

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="C2410C")  # the app's orange

    def style_header(ws, ncols):
        for i in range(1, ncols + 1):
            c = ws.cell(row=1, column=i)
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(vertical="center")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    def fit(ws, ncols):
        for i in range(1, ncols + 1):
            longest = max((len(str(c.value or ""))
                           for c in ws[get_column_letter(i)]), default=0)
            ws.column_dimensions[get_column_letter(i)].width =                 min(52, max(12, longest + 2))

    wb = Workbook()
    ws = wb.active
    ws.title = "Contacts"
    ws.append(headers)
    for r in rows:
        ws.append(r)
    style_header(ws, width)
    fit(ws, width)

    # ── flat one-address-per-row sheet ──
    sc = body.store_col if 0 <= body.store_col < width else -1
    flat = []
    for r in rows:
        store = r[sc].strip() if sc >= 0 else ""
        for i, cell in enumerate(r):
            if i == sc:
                continue
            raw = cell.strip()
            if not raw:
                continue
            em = slc._label_to_email(raw) or (raw if "@" in raw else "")
            if not em:
                continue
            name = raw.split("<")[0].strip() if "<" in raw else ""
            flat.append([store, name, em])

    if flat:
        ws2 = wb.create_sheet("Emails")
        ws2.append(["Store / Vendor", "Name", "Email"])
        for r in flat:
            ws2.append(r)
        style_header(ws2, 3)
        fit(ws2, 3)

    buf = io.BytesIO()
    wb.save(buf)
    name = f"vendor_contacts_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type=("application/vnd.openxmlformats-officedocument"
                    ".spreadsheetml.sheet"),
        headers={"Content-Disposition": f'attachment; filename="{name}"'})


# ── 🧾 missing-item pick sheet ────────────────────────────────
@app.post("/api/picksheet/build")
def picksheet_build(body: PickSheetBody):
    """Selected invoices -> shortfall rows (JSON preview for the panel)."""
    if not body.digits:
        raise HTTPException(400, "Pick at least one invoice.")
    try:
        return picksheet.build(body.digits)
    finally:
        deps.progress(active=False)


@app.post("/api/picksheet/xlsx")
def picksheet_xlsx(body: PickSheetBody):
    if not body.digits:
        raise HTTPException(400, "Pick at least one invoice.")
    try:
        data = picksheet.build(body.digits)
        blob = picksheet.to_xlsx(data)
        tasklog.log_many("picksheet_xlsx", body.digits)   # +1 per invoice
    finally:
        deps.progress(active=False)
    name = f"missing_items_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    return Response(
        content=blob,
        media_type=("application/vnd.openxmlformats-officedocument"
                    ".spreadsheetml.sheet"),
        headers={"Content-Disposition": f'attachment; filename="{name}"'})


@app.post("/api/picksheet/pdf")
def picksheet_pdf(body: PickSheetBody):
    if not body.digits:
        raise HTTPException(400, "Pick at least one invoice.")
    try:
        data = picksheet.build(body.digits)
        blob = picksheet.to_pdf(data)
        tasklog.log_many("picksheet_pdf", body.digits)   # +1 per invoice
    finally:
        deps.progress(active=False)
    name = f"missing_items_{datetime.now():%Y%m%d_%H%M%S}.pdf"
    return Response(content=blob, media_type="application/pdf",
                    headers={"Content-Disposition":
                             f'attachment; filename="{name}"'})


class BulkDateRow(BaseModel):
    digits: str
    delivery_date: str


class BulkDatesBody(BaseModel):
    rows: list[BulkDateRow]


@app.post("/api/orders/bulk-dates")
def orders_bulk_dates(body: BulkDatesBody):
    """Submit the ship-date grid: each row carries its OWN delivery date —
    one verified PATCH per order, exactly like the single-order editor."""
    sess = S.get("bapi_session") or {}
    if not slc._session_ready(sess):
        return {"ok": False, "message": "Needs the 🔑 b-api session — grab it up top."}
    rows = [r for r in body.rows if (r.delivery_date or "").strip()]
    if not rows:
        return {"ok": False, "message": "Give at least one order a date first."}
    oks, fails = [], []
    for i, r in enumerate(rows):
        dg = r.digits
        inv = f"Twiste-{dg}"
        day = str(r.delivery_date)[:10]
        deps.progress(f"🚚 Ship date {inv} ({i + 1}/{len(rows)})…",
                      int(100 * i / len(rows)))
        try:
            od, err = bridge.fetch_dealflow_order(dg)
            if not od:
                fails.append(f"{inv} — {err or 'not found'}")
                continue
            if str(od.get("delivery_date") or "")[:10] == day:
                oks.append(inv)          # already set — nothing to push
                continue
            log = []
            ok, fresh = slc.bapi_update_order_fields(
                sess, od, {"delivery_date": day}, log)
            bridge.push_log(log)
            fresh = fresh or {}
            if ok and str(fresh.get("delivery_date") or "")[:10] not in ("", day):
                fails.append(f"{inv} — Apex kept "
                             f"{str(fresh.get('delivery_date'))[:10]}")
                continue
            if ok:
                if fresh.get("id"):
                    bridge._order_cache_store(dg, fresh, fresh.get("id"))
                oks.append(inv)
                tasklog.log_task("delivery_date_set", invoice=inv, detail=day)
            else:
                fails.append(f"{inv} — Apex rejected the update")
        except Exception as ex:
            fails.append(f"{inv} — {type(ex).__name__}")
    deps.progress(active=False)
    msg = f"🚚 Ship dates: {len(oks)}/{len(rows)} order(s) updated."
    if fails:
        msg += " ❌ " + " · ".join(fails[:5]) + ("…" if len(fails) > 5 else "")
    return {"ok": not fails, "message": msg}


# ── 🚚 delivery instruction grid ─────────────────────────────
# Pick-up date / Layover / Shipper are INTERNAL notes (never pushed to
# Apex) — saved locally like every other cache. Delivery date is the one
# column that PATCHes Apex (via /api/orders/bulk-dates).
DELIVERY_NOTES_FILE = os.path.join(deps.DATA_DIR, "apex_delivery_notes.json")

DELIVERY_EMAIL_TO = [
    "victoria@twistedgrowers.com",
    "kyle@twistedgrowers.com",
    "jacqueline@twistedgrowers.com",
    "bert@twistedgrowers.com",
]
DELIVERY_EMAIL_CC = ["maji@twistedgrowers.com"]


class DeliveryNotesBody(BaseModel):
    notes: dict[str, dict] = {}


@app.get("/api/orders/delivery-notes")
def delivery_notes_get():
    return {"notes": (deps.read_json(DELIVERY_NOTES_FILE, {}) or {})}


_NOTE_KINDS = {"pickup": "pickup_date_set", "layover": "layover_set",
               "shipper": "shipper_set"}


@app.post("/api/orders/delivery-notes")
def delivery_notes_save(body: DeliveryNotesBody):
    """Save the internal grid columns. Scoring: +1 per invoice per field the
    save actually CHANGED to a non-empty value (pickup / layover / shipper
    each their own task) — re-saving the same value earns nothing."""
    changed = []   # (kind, digits, value)

    def _mut(data):
        for dg, n in (body.notes or {}).items():
            if not isinstance(n, dict):
                continue
            ent = data.setdefault(str(dg), {})
            for k in ("pickup", "layover", "shipper"):
                if k not in n:
                    continue
                new_v = str(n.get(k) or "")[:80]
                if new_v and new_v != (ent.get(k) or ""):
                    changed.append((_NOTE_KINDS[k], str(dg), new_v))
                ent[k] = new_v
        return True

    deps.update_json(DELIVERY_NOTES_FILE, {}, _mut)
    for kind, dg, val in changed:
        tasklog.log_task(kind, invoice=f"Twiste-{dg}", detail=val)
    return {"ok": True, "message": (f"💾 Delivery notes saved for "
                                    f"{len(body.notes or {})} order(s)"
                                    + (f" · {len(changed)} field(s) logged"
                                       if changed else "") + ".")}


class DeliveryEmailRow(BaseModel):
    invoice: str
    store: str = ""
    pickup: str = ""
    delivery: str = ""
    layover: str = ""
    shipper: str = ""


class DeliveryEmailBody(BaseModel):
    rows: list[DeliveryEmailRow]
    preview: bool = False
    subject: str = ""
    intro: str = ""
    # None → the coded default lists; an explicit list (even shorter — the
    # preview's ✕ chips) is used exactly as sent
    to: Optional[list[str]] = None
    cc: Optional[list[str]] = None


@app.post("/api/orders/delivery-email")
def delivery_email(body: DeliveryEmailBody):
    """Compose (preview=true → return the default subject/recipients WITHOUT
    sending) or send the delivery instruction grid to the ops list
    (To: Victoria, Kyle, Jacqueline, Bert · Cc: Maji). The sent email is an
    Excel-style table — every cell bordered — with the (editable) intro line
    above it. Subject/intro/rows arrive exactly as edited in the preview."""
    if not body.rows:
        raise HTTPException(400, "Nothing to send — the grid is empty.")

    default_subject = (f"Delivery instructions — {len(body.rows)} order(s) · "
                       f"{datetime.now():%m/%d/%Y}")
    if body.preview:
        return {"ok": True, "to": DELIVERY_EMAIL_TO, "cc": DELIVERY_EMAIL_CC,
                "subject": default_subject,
                "intro": "Delivery instructions:"}

    subject = (body.subject or "").strip() or default_subject
    intro = (body.intro or "").strip()
    to_list = ([a.strip() for a in body.to if a and a.strip()]
               if body.to is not None else list(DELIVERY_EMAIL_TO))
    cc_list = ([a.strip() for a in body.cc if a and a.strip()]
               if body.cc is not None else list(DELIVERY_EMAIL_CC))
    if not to_list:
        raise HTTPException(400, "Keep at least one To recipient.")

    def esc(s):
        return (str(s or "").replace("&", "&amp;").replace("<", "&lt;")
                .replace(">", "&gt;"))

    B = "border:1px solid #000;"           # every cell fully bordered
    th = f"{B}padding:4px 10px;text-align:left;font-weight:bold;background:#f2f2f2"
    td = f"{B}padding:4px 10px"
    rows_html = "".join(
        f"<tr><td style='{td}'>{esc(r.invoice)}</td>"
        f"<td style='{td}'>{esc(r.store)}</td>"
        f"<td style='{td}'>{esc(r.pickup)}</td>"
        f"<td style='{td}'>{esc(r.delivery)}</td>"
        f"<td style='{td}'>{esc(r.layover)}</td>"
        f"<td style='{td}'>{esc(r.shipper)}</td></tr>"
        for r in body.rows)
    html_body = (
        (f"<p>{esc(intro)}</p>" if intro else "")
        + "<table style='border-collapse:collapse;border:1px solid #000;"
          "font-family:Calibri,Arial,sans-serif;font-size:14px'>"
        + f"<tr><th style='{th}'>Invoice</th><th style='{th}'>Store name</th>"
          f"<th style='{th}'>Pick-up date</th><th style='{th}'>Delivery date</th>"
          f"<th style='{th}'>Layover Y/N</th><th style='{th}'>Shipper</th></tr>"
        + rows_html + "</table>")

    ok, msg = slc.send_graph_mail(to_list, subject, html_body,
                                  cc_addrs=cc_list, html=True)
    if ok:
        # +1 for sending the full grid (its own task, not a generic email)
        tasklog.log_task("delivery_grid_sent",
                         detail=f"{len(body.rows)} order(s) · to "
                                f"{len(to_list)} + cc {len(cc_list)}")
    return {"ok": ok, "message": (f"📧 Delivery instructions sent — "
                                  f"{len(body.rows)} order(s) to "
                                  f"{len(to_list)} + cc {len(cc_list)}."
                                  if ok else f"❌ {msg}")}


# ── 📋 FULL picking list (ob-style: brand pages, CS/Units per store) ──
@app.post("/api/pick/full/pdf")
def pick_full_pdf(body: PickSheetBody):
    if not body.digits:
        raise HTTPException(400, "Check at least one order.")
    from . import picking_full
    bundle, invs, missing = picking_full.build_bundle_for(body.digits)
    if bundle is None or bundle.empty:
        raise HTTPException(400, "No items found in the checked orders"
                            + (f" ({', '.join(missing[:5])} unreadable)" if missing else "") + ".")
    blob = picking_full.generate_picking_pdf(bundle, invs)
    # +1 PDF per order that actually made the sheet
    tasklog.log_many("picksheet_pdf",
                     [d for d in body.digits if f"Twiste-{d}" not in missing])
    name = f"picking_list_{datetime.now():%Y%m%d_%H%M%S}.pdf"
    return Response(content=blob, media_type="application/pdf",
                    headers={"Content-Disposition":
                             f'attachment; filename="{name}"'})


@app.post("/api/pick/full/xlsx")
def pick_full_xlsx(body: PickSheetBody):
    if not body.digits:
        raise HTTPException(400, "Check at least one order.")
    from . import picking_full
    bundle, invs, missing = picking_full.build_bundle_for(body.digits)
    if bundle is None or bundle.empty:
        raise HTTPException(400, "No items found in the checked orders"
                            + (f" ({', '.join(missing[:5])} unreadable)" if missing else "") + ".")
    blob = picking_full.generate_picking_excel(bundle)
    # +1 Excel per invoice that actually made the sheet
    tasklog.log_many("picksheet_xlsx",
                     [d for d in body.digits if f"Twiste-{d}" not in missing])
    name = f"picking_list_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    return Response(
        content=blob,
        media_type=("application/vnd.openxmlformats-officedocument"
                    ".spreadsheetml.sheet"),
        headers={"Content-Disposition": f'attachment; filename="{name}"'})


# ── 📊 task activity log ─────────────────────────────────────
class TaskLogBody(BaseModel):
    kind: str
    invoice: str = ""
    detail: str = ""
    items: list[str] = []   # non-empty → one event per item (detail = item)


@app.post("/api/tasklog/log")
def tasklog_ingest(body: TaskLogBody):
    """External ingest — sibling apps (the inventory app's Lab PDF Adder /
    Metrc Tagging) report completed work here so it scores on the dashboard.
    Only known kinds are accepted."""
    if body.kind not in tasklog.KINDS:
        raise HTTPException(400, f"Unknown task kind '{body.kind}'.")
    if body.items:
        tasklog.log_many(body.kind, body.items, invoice=body.invoice)
        n = len(body.items)
    else:
        tasklog.log_task(body.kind, invoice=body.invoice, detail=body.detail)
        n = 1
    return {"ok": True, "message": f"+{n} {tasklog.KINDS[body.kind]}"}


class TaskIdsBody(BaseModel):
    ids: list[str] = []


@app.post("/api/tasklog/delete")
def tasklog_delete(body: TaskIdsBody):
    """Excel-style multi-delete: move the selected events to the deleted bin —
    they drop out of every count but stay restorable."""
    n = tasklog.delete_events(body.ids)
    return {"ok": True, "message": f"🗑 {n} event(s) moved to deleted."}


@app.post("/api/tasklog/restore")
def tasklog_restore(body: TaskIdsBody):
    n = tasklog.restore_events(body.ids)
    return {"ok": True, "message": f"↩ {n} event(s) restored to the log."}


@app.post("/api/tasklog/purge")
def tasklog_purge(body: TaskIdsBody):
    """Permanently clear the deleted bin (all of it when ids is empty)."""
    n = tasklog.purge_deleted(body.ids or None)
    return {"ok": True, "message": f"🧹 {n} deleted event(s) cleared for good."}


@app.get("/api/tasklog/summary")
def tasklog_summary(days: int = 30, start: str | None = None,
                    end: str | None = None, recent_limit: int = 300):
    """Live totals for the Tasks dashboard — per-kind counts (range + today),
    daily/hourly/minute series, and the recent-activity feed (the audit log:
    recent_limit=0 returns EVERY event). start/end (YYYY-MM-DD) are the 📅
    custom calendar range and win over days."""
    return tasklog.summary(days, start=start, end=end,
                           recent_limit=recent_limit)


# ── 📨 email blast (one message per vendor, paced, killable) ─
@app.post("/api/blast/start")
def blast_start(body: BlastBody):
    return blast.start(body.groups, body.subject, body.body, body.cc,
                       body.attachments, body.gap_seconds)


@app.get("/api/blast/status")
def blast_status():
    return blast.status()


# ── 💰 unpaid invoices (Delivered + payment_status unpaid) ───
class UnpaidInvoiceRef(BaseModel):
    id: str
    invoice: str = ""


class UnpaidGroup(BaseModel):
    store: str = ""
    emails: list[str] = []
    invoices: list[UnpaidInvoiceRef] = []


class UnpaidStartBody(BaseModel):
    groups: list[UnpaidGroup]
    subject: str = ""
    body: str = ""
    cc: list[str] = []
    attachments: list[BlastAttachment] = []
    gap_seconds: float = 2.0


@app.get("/api/unpaid/list")
def unpaid_list(live: bool = False):
    """Saved locally like every other cache — the default read serves the
    banked scan instantly; live=true (↻) re-pulls from Apex, re-saves, and
    scores +1 invoice_linked per NEWLY linked invoice."""
    if not live:
        cached = bridge.load_unpaid_cache()
        if cached:
            return {"vendors": cached.get("vendors") or [], "note": "",
                    "at": cached.get("at") or "", "cached": True}
    vendors, note, at = bridge.unpaid_delivered()
    return {"vendors": vendors, "note": note, "at": at, "cached": False}


@app.get("/api/unpaid/invoice-pdf")
def unpaid_invoice_pdf(id: str, invoice: str = ""):
    """Build ONE invoice PDF exactly as the send would and stream it inline —
    the 👁 preview's proof of what actually goes out. Nothing is logged; this
    is verification, not a task."""
    sess = S.get("bapi_session") or {}
    if not slc._session_ready(sess):
        raise HTTPException(502, "Needs the 🔑 b-api session.")
    atts, failed = blast._build_invoice_atts(sess, [{"id": id,
                                                     "invoice": invoice or id}])
    if not atts:
        raise HTTPException(502, f"Couldn't build the PDF for {invoice or id}.")
    name, blob, mime = atts[0]
    return Response(content=blob, media_type=mime or "application/pdf",
                    headers={"Content-Disposition": f'inline; filename="{name}"'})


@app.post("/api/unpaid/start")
def unpaid_start(body: UnpaidStartBody):
    """Collections run: one email per vendor with THEIR unpaid invoice PDFs
    auto-attached (built fresh at send time) + any shared extra files.
    Scoring: invoice_linked is scored at ↻ refresh (the scan); the worker
    logs +1 unpaid_email per send and +1 pdf_attached per attachment."""
    return blast.start(body.groups, body.subject, body.body, body.cc,
                       body.attachments, body.gap_seconds,
                       email_kind="unpaid_email", require_attachments=False)


@app.post("/api/blast/stop")
def blast_stop():
    return blast.stop()


# ── 📦 batch inventory — view / edit (port of the 2.0 grid) ──
@app.get("/api/batches")
def batches_list(name: str = "", hide_archived: bool = True):
    """All batches joined with the catalog. First call triggers the cached
    /v2/batches crawl (~15 min ttl) — can take a while cold."""
    try:
        deps.progress("📦 Loading batch inventory…", active=True)
        return inventory.list_batches(name, hide_archived)
    finally:
        deps.progress(active=False)


@app.post("/api/batches/set")
def batches_set(body: BatchSetBody):
    """Absolute-qty and/or name write for one batch (the 2.0 autosave)."""
    if body.quantity is None and not (body.name or "").strip():
        raise HTTPException(400, "Give a quantity and/or a name.")
    ok, msg = inventory.set_batch(body.batch_id, body.product_id,
                                  body.quantity, body.name)
    if not ok:
        raise HTTPException(502, msg)
    if body.quantity is not None:
        tasklog.log_task("inventory_adjustment",
                         detail=f"batch {body.batch_id} → {body.quantity}")
    return {"ok": True, "message": msg}


# ── 🤖 AI order builder (local Ollama, read-only preview) ────
@app.post("/api/ai/parse-order")
def ai_parse_order(body: AiParseBody):
    """Free text -> catalog-matched line items. Preview only — writes nothing."""
    if not (body.text or "").strip():
        raise HTTPException(400, "Type an order first.")
    try:
        return ai.parse_order(body.text)
    except ai.OllamaError as e:
        raise HTTPException(503, str(e))
    finally:
        deps.progress(active=False)


@app.post("/api/ai/chat")
def ai_chat(body: AiChatBody):
    """Bridge Assistant turn. The model can read (board/catalog/contacts)
    and stage email drafts — sending only ever happens via /api/ai/send-email
    below, triggered by the human clicking Send on a staged draft."""
    if not any((m.get("content") or "").strip() for m in body.messages):
        raise HTTPException(400, "Say something first.")
    try:
        return ai.chat(body.messages)
    except ai.OllamaError as e:
        raise HTTPException(503, str(e))
    finally:
        deps.progress(active=False)


@app.post("/api/ai/send-email")
def ai_send_email(body: AiSendBody):
    ok, msg = ai.send_email(body.to, body.cc, body.subject, body.body)
    if not ok:
        raise HTTPException(502, msg)
    return {"ok": True, "message": msg}


@app.post("/api/ai/learn-style")
def ai_learn_style():
    """Read recent Sent Items and distill Vincent's writing style locally."""
    try:
        return ai.learn_style()
    except ai.OllamaError as e:
        raise HTTPException(503, str(e))
    finally:
        deps.progress(active=False)


# ── 📧 email + downloads ──────────────────────────────────────
@app.get("/api/email/prepare")
def email_prepare():
    return review.email_prepare()


@app.post("/api/email/send")
def email_send(body: EmailBody):
    try:
        ok, msg = review.send_email(body.to, body.cc, body.subject, body.body,
                                    body.attach, extra=body.extra)
    finally:
        deps.progress(active=False)
    return {"ok": ok, "message": (f"✅ {msg}" if ok else f"❌ {msg}")}


@app.get("/api/download/invoice")
def download_invoice():
    blob = review._invoice_blob()
    return Response(content=blob["data"],
                    media_type=blob.get("mime", "application/pdf"),
                    headers={"Content-Disposition":
                             f'attachment; filename="{blob["name"]}"'})


@app.get("/api/download/bundle")
def download_bundle():
    data, name, _msg = review.build_bundle()
    return Response(content=data, media_type="application/zip",
                    headers={"Content-Disposition":
                             f'attachment; filename="{name}"'})


@app.get("/api/download/coa")
def download_coa():
    data, name, msg = review.coa_zip()
    if not data:
        raise HTTPException(502, msg)
    return Response(content=data, media_type="application/zip",
                    headers={"Content-Disposition":
                             f'attachment; filename="{name}"'})


# ── 🧪 Lab PDF Batch Adder (native port of the inventory app) ─
@app.get("/api/lab/meta")
def lab_meta(with_ops: bool = True):
    """Constant tables + operations for the /lab page."""
    ops = labtools.fetch_operations() if with_ops else {}
    cache = metrctools.load_cache() or {}
    return {
        "brands": labtools.BRANDS,
        "categories": labtools.PRODUCT_CATEGORIES,
        "product_types": labtools.PRODUCT_TYPES,
        "case_sizes": {f"{b}|{c}": n
                       for (b, c), n in labtools.DEFAULT_CASE_SIZES.items()},
        "unit_pricing": {f"{b}|{t}": p
                         for (b, t), p in labtools.UNIT_PRICING.items()},
        "unit_sizes": labtools.UNIT_SIZES,
        "operations": ops,
        "default_operation_id": labtools.default_operation_id(ops),
        "products_count": len(cache.get("products") or []),
        "products_at": cache.get("last_updated"),
        "pdf_parser": labtools.PDF_PARSER_AVAILABLE or labtools.PYPDF_AVAILABLE,
    }


@app.post("/api/lab/parse")
async def lab_parse(files: list[UploadFile] = File(...),
                    brand: str = Form("Twisted Buds"),
                    category: str = Form("Pre-pack")):
    """Upload + parse COA PDFs; each is matched to existing products
    (strict brand+category). Files are kept server-side for process time."""
    items = []
    for f in files:
        data = await f.read()
        items.append(labtools.parse_upload(f.filename or "COA.pdf", data,
                                           brand, category))
    return {"ok": True, "items": items}


@app.get("/api/lab/rematch")
def lab_rematch(strain: str, brand: str, category: str):
    return {"matches": labtools.rematch(strain, brand, category)}


@app.get("/api/lab/search")
def lab_search(q: str):
    return {"matches": labtools.search_products(q)}


@app.post("/api/lab/refresh-products")
def lab_refresh_products():
    products = labtools.get_products(force_fetch=True)
    return {"ok": True, "message": f"Pulled {len(products)} products from Apex."}


class LabProcessBody(BaseModel):
    items: list[dict] = []


@app.post("/api/lab/process")
def lab_process(body: LabProcessBody):
    """Create All Batches — per item: product (optional) → batch → COA upload →
    terpenes. Scoring: +1 per line box filled, +1 per COA attached."""
    results = []
    n = len(body.items)
    for i, item in enumerate(body.items):
        deps.progress(f"Creating batch {i + 1}/{n}…", int(100 * i / max(1, n)))
        try:
            results.append(labtools.process_item(item))
        except Exception as e:
            results.append({"ok": False, "error": str(e)})
    deps.progress(active=False)
    done = sum(1 for r in results if r.get("ok"))
    return {"ok": True, "results": results,
            "message": f"✅ {done}/{n} batch(es) created."}


# ── 🏷️ Metrc Tagging (native port of the inventory app) ─
@app.get("/api/mtag/state")
def mtag_state():
    """Grid rows + cache status for the /metrc page. Local caches only — Apex
    is touched exclusively by the pull endpoints."""
    metrctools.seed_from_inventory_app()
    cache = metrctools.load_cache()
    posted = metrctools.load_posted_tags()
    linked = metrctools.load_linked_tags()
    rows = metrctools.build_inventory_rows(cache, posted, linked) if cache else []
    pkgs = metrctools.load_packages_cache() or {}
    return {
        "rows": rows,
        "last_updated": (cache or {}).get("last_updated"),
        "products": len((cache or {}).get("products") or []),
        "batches": len((cache or {}).get("batches") or []),
        "packages": len(pkgs.get("packages_by_id") or {}),
        "packages_at": pkgs.get("last_updated"),
        "linked_count": len(linked),
        "session_ready": metrctools.session_ready(),
    }


@app.post("/api/mtag/fresh-pull")
def mtag_fresh_pull():
    """Bulk re-pull of products + batches (the only metered call), then — if the
    b-api session is live — re-harvest existing Metrc tags."""
    try:
        cache = metrctools.fresh_pull()
        msg = (f"Pulled {len(cache.get('products') or [])} products, "
               f"{len(cache.get('batches') or [])} batches.")
        if metrctools.session_ready():
            deps.progress("Reading existing Metrc tags…", 50)
            mapping, err = metrctools.bapi_harvest_linked_tags()
            if mapping:
                metrctools.save_linked_tags(mapping)
                msg += f" Found existing tags on {len(mapping)} batches."
        return {"ok": True, "message": msg}
    finally:
        deps.progress(active=False)


@app.post("/api/mtag/pull-packages")
def mtag_pull_packages():
    """Pull the Metrc source-package options (b-api, per cached operation)."""
    if not metrctools.session_ready():
        raise HTTPException(400, "Apex session not set — grab it on the main "
                                 "page first (🔐 Session).")
    cache = metrctools.load_cache() or {}
    op_ids = sorted({b.get("operation_id") for b in cache.get("batches", [])
                     if isinstance(b, dict) and b.get("operation_id")})
    if not op_ids:
        raise HTTPException(400, "No operation ids in the cached batches — do a "
                                 "Fresh Pull first.")
    by_id, errs = metrctools.bapi_fetch_packages(op_ids)
    if by_id:
        metrctools.save_packages_cache({
            "last_updated": datetime.now().isoformat(timespec="seconds"),
            "packages_by_id": {str(k): v for k, v in by_id.items()},
        })
    msg = f"Pulled {len(by_id)} Metrc packages."
    if errs:
        msg += " Issues: " + " · ".join(errs[:3])
    return {"ok": bool(by_id), "message": msg}


@app.post("/api/mtag/pull-tags")
def mtag_pull_tags():
    """Harvest every batch's existing linked Metrc tag from Apex inventory."""
    if not metrctools.session_ready():
        raise HTTPException(400, "Apex session not set — grab it on the main "
                                 "page first (🔐 Session).")
    try:
        mapping, err = metrctools.bapi_harvest_linked_tags()
    finally:
        deps.progress(active=False)
    if not mapping:
        raise HTTPException(502, f"Couldn't read linked tags: {err}")
    metrctools.save_linked_tags(mapping)
    return {"ok": True,
            "message": f"Found existing Metrc tags on {len(mapping)} batches."}


@app.get("/api/mtag/packages")
def mtag_packages():
    pkgs = metrctools.load_packages_cache() or {}
    return {"packages": list((pkgs.get("packages_by_id") or {}).values()),
            "last_updated": pkgs.get("last_updated")}


class MtagSaveBody(BaseModel):
    changes: list[dict] = []   # {batch_id, name?, quantity?}
    tags: list[dict] = []      # {batch_id, package_id, label?}


@app.post("/api/mtag/save")
def mtag_save(body: MtagSaveBody):
    """Save changed rows to Apex: name/qty edits PATCH the public v2 batch;
    tag picks post via the b-api read-modify-write. +1 metrc_tagged per
    successfully posted tag."""
    ok = fail = tagged = 0
    errors = []
    cache = metrctools.load_cache()

    for ch in body.changes:
        bid = ch.get("batch_id")
        updates = {k: v for k, v in ch.items()
                   if k in ("name", "quantity") and v is not None}
        if not bid or not updates:
            continue
        success, msg = metrctools.update_batch(bid, updates)
        if success:
            ok += 1
            for b in (cache or {}).get("batches", []):
                if b.get("id") == bid:
                    b.update(updates)
                    break
        else:
            fail += 1
            errors.append(f"batch {bid}: {msg}")
    if ok and cache:
        metrctools.save_cache(cache)

    if body.tags:
        if not metrctools.session_ready():
            raise HTTPException(400, "Tag(s) not posted: Apex session not set.")
        pkgs = (metrctools.load_packages_cache() or {}).get("packages_by_id") or {}
        posted = metrctools.load_posted_tags()
        linked = metrctools.load_linked_tags()
        n = len(body.tags)
        for i, t in enumerate(body.tags):
            bid = t.get("batch_id")
            pkg = pkgs.get(str(t.get("package_id")))
            if not bid or not pkg:
                fail += 1
                errors.append(f"tag for batch {bid}: package not in cache")
                continue
            deps.progress(f"Posting tag {i + 1}/{n}…", int(100 * i / n))
            success, msg, _ = metrctools.bapi_post_tag(bid, [pkg])
            if success:
                tagged += 1
                tag_val = pkg.get("Label", "")
                posted[str(bid)] = tag_val
                linked[str(bid)] = tag_val
                # 📊 +1 per METRC item tagged
                tasklog.log_task('metrc_tagged',
                                 detail=f"{tag_val} · batch #{bid}")
            else:
                fail += 1
                errors.append(f"tag batch {bid}: {msg}")
        metrctools.save_posted_tags(posted)
        metrctools.save_linked_tags(linked)
        deps.progress(active=False)

    parts = []
    if ok:
        parts.append(f"✅ {ok} batch change(s) saved")
    if tagged:
        parts.append(f"🏷️ {tagged} Metrc tag(s) posted")
    if fail:
        parts.append(f"❌ {fail} failed")
    return {"ok": fail == 0, "message": " · ".join(parts) or "No changes.",
            "errors": errors[:10]}


@app.get("/api/health")
def health():
    return {"ok": True, "engine": "slc.py (copied)", "data_dir": deps.DATA_DIR}
