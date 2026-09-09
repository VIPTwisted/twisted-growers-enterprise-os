"""Full picking list — ported from ob (app21.py)'s picking bundle builders so
the deal-flow section generates the SAME pick sheets as the Apex Picking
Manager: per-brand landscape PDF pages with tick boxes, and an Excel workbook
with a Master sheet + one sheet per brand, quantities split CS / Units per
store. Fed from checked deal-flow orders via the order cache (0 metered)."""

import re
from datetime import datetime
from io import BytesIO

import pandas as pd

from . import bridge


# ── item parsing (verbatim from ob/backend/engine/app21.py) ──

def extract_batch(item: dict) -> str:
    candidate_keys = (
        "batch_name", "batch_number", "batch",
        "lot_number", "lot_name", "lot",
    )
    for key in candidate_keys:
        val = item.get(key)
        if isinstance(val, str) and val.strip():
            return val.strip()
        if isinstance(val, dict):
            name = val.get("name") or val.get("batch_name") or val.get("number")
            if isinstance(name, str) and name.strip():
                return name.strip()
    return ""


def build_product_display(product_name: str, batch: str) -> str:
    if batch:
        return f"{product_name} - Batch: {batch}"
    return product_name


def extract_unit_type(item: dict) -> str:
    """"CS" or "Units" off the ORDER-level unit-of-measure (same rules as ob —
    the base `unit_measurement` field is deliberately ignored)."""
    oum = item.get("order_unit_measurement")
    if isinstance(oum, dict):
        name = (oum.get("name") or "").strip().lower()
        alias = (oum.get("alias") or "").strip().lower()
        if name == "case" or alias == "cs":
            return "CS"
        if name == "unit" or alias in ("u", "ea", "each"):
            return "Units"
    oum_id = item.get("order_unit_measurement_id")
    if oum_id == 21:
        return "CS"
    if oum_id == 14:
        return "Units"
    string_keys = (
        "order_unit", "order_unit_type", "uom",
        "sold_by", "sold_as", "pricing_unit", "quantity_type", "quantity_unit",
    )
    for key in string_keys:
        val = item.get(key)
        if isinstance(val, str) and val.strip():
            v = val.strip().lower()
            if "case" in v or v in ("cs", "cases"):
                return "CS"
            if "unit" in v or v in ("ea", "each", "units"):
                return "Units"
        if isinstance(val, dict):
            v = (val.get("name") or val.get("alias") or "").strip().lower()
            if v in ("case", "cs"):
                return "CS"
            if v in ("unit", "u", "ea", "each"):
                return "Units"
    for key in ("sold_by_unit", "sell_by_unit", "by_unit", "is_unit"):
        if item.get(key) is True:
            return "Units"
    for key in ("is_case", "by_case", "sold_by_case"):
        if item.get(key) is True:
            return "CS"
    return "CS"


def fmt_qty(case_qty: int, unit_qty: int) -> str:
    parts = []
    if case_qty:
        parts.append(f"{case_qty} CS")
    if unit_qty:
        parts.append(f"{unit_qty} Units")
    return ", ".join(parts)


def split_qty(items_subset) -> tuple:
    cs = int(items_subset[items_subset["unit_type"] == "CS"]["quantity"].sum())
    un = int(items_subset[items_subset["unit_type"] == "Units"]["quantity"].sum())
    return cs, un


def create_picking_bundle(selected_df):
    all_items = []
    for _, row in selected_df.iterrows():
        items = row.get("items", [])
        if not items:
            continue
        for item in items:
            if not isinstance(item, dict):
                continue
            brand = item.get("brand", {})
            brand_name = (brand.get("name", "Unknown Brand")
                          if isinstance(brand, dict) else "Unknown Brand")
            product_name = item.get("product_name", "Unknown Product")
            batch = extract_batch(item)
            product_display = build_product_display(product_name, batch)
            unit_type = extract_unit_type(item)
            quantity = item.get("order_quantity", 0)
            price = item.get("order_price", "0")
            if isinstance(price, str):
                price = float(price.replace(",", "") or 0)
            all_items.append({
                "brand": brand_name,
                "product_name": product_name,
                "batch": batch,
                "product_display": product_display,
                "unit_type": unit_type,
                "quantity": quantity,
                "price": price,
                "invoice_number": row["invoice_number"],
                "buyer": row["buyer_name"],
            })
    if not all_items:
        return pd.DataFrame()
    return pd.DataFrame(all_items).sort_values(["brand", "product_display"])


# ── tg-side bundle builder: checked digits → orders → bundle ──

def build_bundle_for(digits_list):
    """(bundle_df, invoices, missing) for the checked orders — b-api reads by
    the banked order id first (works for orders NOT on the board), then the
    bridge's normal order fetch as fallback."""
    rows, missing = [], []
    for dg in digits_list:
        inv = f"Twiste-{dg}"
        od, _err = bridge.fetch_dealflow_order(dg)
        if not od or not (od.get("items") or []):
            missing.append(inv)
            continue
        buyer = od.get("buyer") or {}
        rows.append({
            "invoice_number": (od.get("invoiceNumber")
                               or od.get("custom_invoice_number")
                               or od.get("invoice_number") or inv),
            "buyer_name": (buyer.get("name", "?")
                           if isinstance(buyer, dict) else "?"),
            "items": od.get("items") or [],
        })
    if not rows:
        return pd.DataFrame(), [], missing
    df = pd.DataFrame(rows)
    return create_picking_bundle(df), [r["invoice_number"] for r in rows], missing


# ── Excel (verbatim from ob) ──

def _safe_sheet_name(name: str, used: set) -> str:
    invalid = set(':\\/?*[]')
    clean = "".join(c for c in str(name) if c not in invalid).strip() or "Sheet"
    clean = clean[:31]
    base = clean
    i = 1
    while clean.lower() in used:
        suffix = f"_{i}"
        clean = base[:31 - len(suffix)] + suffix
        i += 1
    used.add(clean.lower())
    return clean


def generate_picking_excel(bundle_df) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F2937")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    total_font = Font(bold=True)

    def _write_sheet(ws, headers, rows, totals_row, text_cols):
        for c_idx, head in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=c_idx, value=head)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = border
        for r_offset, row_vals in enumerate(rows, start=2):
            for c_idx, val in enumerate(row_vals, start=1):
                cell = ws.cell(row=r_offset, column=c_idx,
                               value=(val if val not in (0, "") else None))
                cell.border = border
                cell.alignment = left if c_idx <= text_cols else center
        total_r = len(rows) + 2
        for c_idx, val in enumerate(totals_row, start=1):
            cell = ws.cell(row=total_r, column=c_idx,
                           value=(val if val not in (0, "") else None))
            cell.border = border
            cell.font = total_font
            cell.alignment = left if c_idx <= text_cols else center
        PAD = 4
        all_rows = rows + [totals_row]
        for c_idx, head in enumerate(headers, start=1):
            letter = ws.cell(row=1, column=c_idx).column_letter
            body_max = max(
                (len(str(r[c_idx - 1])) for r in all_rows
                 if r[c_idx - 1] not in (None, "")),
                default=0,
            )
            if c_idx <= text_cols:
                content = max(len(str(head)), body_max, len("TOTAL"))
                ws.column_dimensions[letter].width = min(content + PAD + 2, 120)
            else:
                longest_word = max((len(w) for w in str(head).split()),
                                   default=len(str(head)))
                content = max(longest_word, body_max, len("TOTAL"))
                ws.column_dimensions[letter].width = min(max(content + PAD, 12), 30)
        ws.freeze_panes = "A2"

    wb = Workbook()
    used_names = set()
    all_stores = sorted(bundle_df["buyer"].unique())
    brands = sorted(bundle_df["brand"].unique())

    master = wb.active
    master.title = _safe_sheet_name("Master", used_names)
    master_headers = ["Brand", "Product"] + all_stores + ["TOTAL"]
    master_rows = []
    col_cs = {s: 0 for s in all_stores}
    col_un = {s: 0 for s in all_stores}
    grand_cs = grand_un = 0
    for brand in brands:
        brand_items = bundle_df[bundle_df["brand"] == brand]
        for product in sorted(brand_items["product_display"].unique()):
            prod_items = brand_items[brand_items["product_display"] == product]
            row = [brand, product]
            row_cs = row_un = 0
            for store in all_stores:
                cs, un = split_qty(prod_items[prod_items["buyer"] == store])
                row.append(fmt_qty(cs, un) or None)
                col_cs[store] += cs
                col_un[store] += un
                row_cs += cs
                row_un += un
            row.append(fmt_qty(row_cs, row_un) or None)
            grand_cs += row_cs
            grand_un += row_un
            master_rows.append(row)
    master_totals = (["TOTAL", None]
                     + [fmt_qty(col_cs[s], col_un[s]) or None for s in all_stores]
                     + [fmt_qty(grand_cs, grand_un) or None])
    _write_sheet(master, master_headers, master_rows, master_totals, text_cols=2)

    for brand in brands:
        brand_items = bundle_df[bundle_df["brand"] == brand]
        brand_stores = sorted(brand_items["buyer"].unique())
        headers = ["Product"] + brand_stores + ["TOTAL"]
        rows = []
        b_col_cs = {s: 0 for s in brand_stores}
        b_col_un = {s: 0 for s in brand_stores}
        b_grand_cs = b_grand_un = 0
        for product in sorted(brand_items["product_display"].unique()):
            prod_items = brand_items[brand_items["product_display"] == product]
            row = [product]
            row_cs = row_un = 0
            for store in brand_stores:
                cs, un = split_qty(prod_items[prod_items["buyer"] == store])
                row.append(fmt_qty(cs, un) or None)
                b_col_cs[store] += cs
                b_col_un[store] += un
                row_cs += cs
                row_un += un
            row.append(fmt_qty(row_cs, row_un) or None)
            b_grand_cs += row_cs
            b_grand_un += row_un
            rows.append(row)
        totals = (["TOTAL"]
                  + [fmt_qty(b_col_cs[s], b_col_un[s]) or None for s in brand_stores]
                  + [fmt_qty(b_grand_cs, b_grand_un) or None])
        ws = wb.create_sheet(title=_safe_sheet_name(brand, used_names))
        _write_sheet(ws, headers, rows, totals, text_cols=1)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ── PDF (verbatim from ob) ──

def generate_picking_pdf(bundle_df, selected_invoices) -> bytes:
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT
    from reportlab.platypus import (
        SimpleDocTemplate, Table, TableStyle, Paragraph, PageBreak, Flowable,
    )

    GRID = colors.HexColor("#BBBBBB")
    HEAD_BG = colors.HexColor("#1F2937")
    TOTAL_BG = colors.HexColor("#E5E7EB")
    ALT_BG = colors.HexColor("#FBFCFD")
    INK = colors.HexColor("#111827")
    MUTED = colors.HexColor("#6B7280")

    name_style = ParagraphStyle("pname", fontName="Helvetica-Bold", fontSize=8.5,
                                leading=10.5, textColor=INK)
    batch_style = ParagraphStyle("pbatch", fontName="Helvetica", fontSize=7,
                                 leading=8.5, textColor=MUTED, spaceBefore=2)
    hdr_l = ParagraphStyle("hdrl", fontName="Helvetica-Bold", fontSize=8,
                           leading=9.5, textColor=colors.white, alignment=TA_LEFT)
    hdr_c = ParagraphStyle("hdrc", fontName="Helvetica-Bold", fontSize=8,
                           leading=9.5, textColor=colors.white, alignment=TA_CENTER)
    cell_c = ParagraphStyle("cellc", fontName="Helvetica", fontSize=8.5,
                            leading=10, alignment=TA_CENTER, textColor=INK)
    cell_c_b = ParagraphStyle("cellcb", fontName="Helvetica-Bold", fontSize=8.5,
                              leading=10, alignment=TA_CENTER, textColor=INK)
    cell_muted = ParagraphStyle("cellm", fontName="Helvetica", fontSize=8.5,
                                leading=10, alignment=TA_CENTER,
                                textColor=colors.HexColor("#C7CDD6"))
    total_lbl = ParagraphStyle("tlbl", fontName="Helvetica-Bold", fontSize=8.5,
                               leading=10, alignment=TA_LEFT, textColor=INK)
    title_style = ParagraphStyle("btitle", fontName="Helvetica-Bold", fontSize=14,
                                 leading=17, textColor=HEAD_BG)

    class _CheckBox(Flowable):
        def __init__(self, size=11):
            Flowable.__init__(self)
            self.size = self.width = self.height = size

        def draw(self):
            c = self.canv
            c.setLineWidth(1.1)
            c.setStrokeColor(colors.HexColor("#9CA3AF"))
            c.roundRect(0, 0, self.size, self.size, 2.2, stroke=1, fill=0)

    class _Tick(Flowable):
        def __init__(self, size=10, color=colors.white):
            Flowable.__init__(self)
            self.size = self.width = self.height = size
            self.color = color

        def draw(self):
            c = self.canv
            s = self.size
            c.setStrokeColor(self.color)
            c.setLineWidth(1.6)
            c.setLineCap(1)
            c.line(0.15 * s, 0.50 * s, 0.42 * s, 0.22 * s)
            c.line(0.42 * s, 0.22 * s, 0.85 * s, 0.78 * s)

    buf = BytesIO()
    PAGE = landscape(letter)
    LM = RM = 30
    doc = SimpleDocTemplate(buf, pagesize=PAGE, leftMargin=LM, rightMargin=RM,
                            topMargin=46, bottomMargin=30, title="Picking List")
    avail = PAGE[0] - LM - RM

    def _clean_name(name: str) -> str:
        return re.sub(r"\s*\(\s*sold[^)]*\)", "", str(name), flags=re.I).strip()

    orders = ", ".join(selected_invoices) if selected_invoices else ""

    def _page_meta(canvas, _doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(MUTED)
        x = PAGE[0] - RM
        y = PAGE[1] - 26
        if orders:
            canvas.drawRightString(x, y, f"Orders: {orders}")
        canvas.drawRightString(x, y - 11,
                               datetime.now().strftime("%B %d, %Y  %I:%M %p"))
        canvas.restoreState()

    elements = []
    brands = sorted(bundle_df["brand"].unique())

    for bi, brand in enumerate(brands):
        bitems = bundle_df[bundle_df["brand"] == brand]
        stores = sorted(bitems["buyer"].unique())
        ncols = len(stores) + 3

        check_w = 28
        total_w = 54
        product_w = avail * 0.30
        store_w = (avail - product_w - total_w - check_w) / max(len(stores), 1)
        col_widths = [product_w] + [store_w] * len(stores) + [total_w, check_w]

        brand_row = [Paragraph(brand, title_style)] + [""] * (ncols - 1)
        header = [Paragraph("Product", hdr_l)]
        for s in stores:
            header.append(Paragraph(s, hdr_c))
        header.append(Paragraph("TOTAL", hdr_c))
        header.append(_Tick())
        data = [brand_row, header]

        store_totals_cs = {s: 0 for s in stores}
        store_totals_un = {s: 0 for s in stores}
        brand_cs = brand_un = 0

        for product in sorted(bitems["product_display"].unique()):
            pit = bitems[bitems["product_display"] == product]
            p_name = _clean_name(pit["product_name"].iloc[0])
            p_batch = pit["batch"].iloc[0] if "batch" in pit else ""

            prod_cell = [Paragraph(p_name, name_style)]
            if p_batch:
                prod_cell.append(Paragraph(f"Batch: {p_batch}", batch_style))

            row = [prod_cell]
            row_cs = row_un = 0
            for s in stores:
                cs, un = split_qty(pit[pit["buyer"] == s])
                store_totals_cs[s] += cs
                store_totals_un[s] += un
                row_cs += cs
                row_un += un
                txt = fmt_qty(cs, un)
                row.append(Paragraph(txt, cell_c) if txt
                           else Paragraph("&mdash;", cell_muted))
            brand_cs += row_cs
            brand_un += row_un
            row.append(Paragraph(fmt_qty(row_cs, row_un), cell_c_b))
            row.append(_CheckBox())
            data.append(row)

        totals = [Paragraph("TOTAL", total_lbl)]
        for s in stores:
            txt = fmt_qty(store_totals_cs[s], store_totals_un[s])
            totals.append(Paragraph(txt or "", cell_c_b))
        totals.append(Paragraph(fmt_qty(brand_cs, brand_un), cell_c_b))
        totals.append(Paragraph("", cell_c))
        data.append(totals)

        last = len(data) - 1
        tbl = Table(data, colWidths=col_widths, repeatRows=2)
        style = [
            ("GRID", (0, 0), (-1, -1), 0.5, GRID),
            ("SPAN", (0, 0), (-1, 0)),
            ("BACKGROUND", (0, 0), (-1, 0), TOTAL_BG),
            ("BACKGROUND", (0, 1), (-1, 1), HEAD_BG),
            ("TOPPADDING", (0, 0), (-1, 0), 6),
            ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
            ("LEFTPADDING", (0, 0), (0, 0), 8),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (-1, 1), (-1, -1), "CENTER"),
            ("LEFTPADDING", (0, 1), (-1, -1), 5),
            ("RIGHTPADDING", (0, 1), (-1, -1), 5),
            ("TOPPADDING", (0, 1), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 1), (-1, -1), 4),
            ("BACKGROUND", (0, last), (-1, last), TOTAL_BG),
        ]
        for r in range(2, last):
            if r % 2 == 1:
                style.append(("BACKGROUND", (0, r), (-1, r), ALT_BG))
        tbl.setStyle(TableStyle(style))
        elements.append(tbl)

        if bi != len(brands) - 1:
            elements.append(PageBreak())

    doc.build(elements, onFirstPage=_page_meta, onLaterPages=_page_meta)
    return buf.getvalue()
